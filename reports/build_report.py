"""Build the project report from verified project artifacts."""
import csv
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.cache' / 'report-tools'))
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
FIG = ROOT / 'figures'
ASSETS = OUT / 'assets'
meta = json.loads((ROOT / 'models' / 'metadata.json').read_text(encoding='utf-8'))
audit = json.loads((ROOT / 'results' / 'data_audit.json').read_text(encoding='utf-8'))
split = list(csv.DictReader((ROOT / 'results' / 'split_summary.csv').open(encoding='utf-8')))
comparison = list(csv.DictReader((ROOT / 'results' / 'model_comparison.csv').open(encoding='utf-8')))
tuning = list(csv.DictReader((ROOT / 'results' / 'tuning_summary.csv').open(encoding='utf-8')))
feature = list(csv.DictReader((ROOT / 'results' / 'feature_selection.csv').open(encoding='utf-8')))
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)
sec.top_margin, sec.bottom_margin = Inches(.75), Inches(.72)
sec.left_margin, sec.right_margin = Inches(1.15), Inches(.85)
sec.header_distance, sec.footer_distance = Inches(.32), Inches(.35)
navy = RGBColor(22, 52, 62)
teal = RGBColor(23, 119, 112)
muted = RGBColor(96, 111, 119)

normal = doc.styles['Normal']
normal.font.name, normal.font.size, normal.font.color.rgb = 'Times New Roman', Pt(11.5), navy
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.22
for name, size, before, after in [('Title', 22, 0, 16), ('Heading 1', 16, 14, 8), ('Heading 2', 13, 11, 5), ('Heading 3', 11.5, 8, 4)]:
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.bold, st.font.color.rgb = 'Cambria', Pt(size), True, navy
    st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(after)
    st.paragraph_format.keep_with_next = True
for name, size in [('Caption', 9), ('Small', 9.5)]:
    if name not in doc.styles:
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st = doc.styles[name]
    st.font.name, st.font.size, st.font.color.rgb = 'Times New Roman', Pt(size), muted
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = name == 'Caption'

def field(p, instruction):
    r = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    default = OxmlElement('w:t'); default.text = 'Update field in Word'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, separate, default, end): r._r.append(el)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.style = 'Small'
header.add_run('CREDIT CARD FRAUD DETECTION  |  PROJECT REPORT')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.style = 'Small'
footer.add_run('Credit Card Fraud Detection  •  ')
field(footer, 'PAGE')

def para(text='', style=None, align=None, bold=False):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    return p

def heading(text, level=1):
    doc.add_heading(text, level)

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    r.font.name = 'Times New Roman'; r.font.size = Pt(9)
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs: r.bold = True; r.font.color.rgb = navy
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

figure_index = []
def image(path, caption, width=6.05):
    path = Path(path)
    if not path.exists(): return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = para(f'Figure {len(figure_index)+1}. {caption}', 'Caption', WD_ALIGN_PARAGRAPH.CENTER)
    figure_index.append(caption)

def newpage(): doc.add_page_break()
def chapter(n, title):
    newpage(); heading(f'CHAPTER {n}  |  {title.upper()}', 1)

def blank_line(label):
    para(f'{label}:  ' + '_' * 43)

def metric(v, digits=4): return f'{float(v):.{digits}f}'

# Front matter. Identity and institutional fields are intentionally blank.
para('PROJECT REPORT', 'Title', WD_ALIGN_PARAGRAPH.CENTER)
para('CREDIT CARD FRAUD DETECTION', 'Title', WD_ALIGN_PARAGRAPH.CENTER)
para('Using Machine Learning', align=WD_ALIGN_PARAGRAPH.CENTER)
para('Fraud Signal Studio: a Streamlit demonstration', align=WD_ALIGN_PARAGRAPH.CENTER)
para('\nSubmitted in partial fulfillment of the requirements for', align=WD_ALIGN_PARAGRAPH.CENTER)
blank_line('Degree / course')
blank_line('Department')
blank_line('Institution')
blank_line('University')
blank_line('Academic year')
para('\nSubmitted by', align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
table(['No.', 'Student name', 'Roll / enrollment number'], [[i, '____________________________', '________________________'] for i in range(1, 5)])
blank_line('Project guide')
blank_line('Group / batch number')
para('\nSubmission date:  ____________________', align=WD_ALIGN_PARAGRAPH.CENTER)

newpage(); heading('CERTIFICATE', 1)
para('This is to certify that the project report entitled “Credit Card Fraud Detection Using Machine Learning” has been prepared and submitted by the students listed below in partial fulfillment of the requirements of the degree / course stated on the title page. The work described in this report concerns the ULB / Worldline credit-card transaction benchmark and its machine learning classification workflow.')
table(['No.', 'Student name', 'Roll / enrollment number'], [[i, '____________________________', '________________________'] for i in range(1, 5)])
para('The students completed the project under the supervision of the guide named below. The institution may complete and certify this page after reviewing the report and project materials.')
blank_line('Guide name and signature')
blank_line('Head of department name and signature')
blank_line('Principal / authorized signatory')
blank_line('Institution / department')
blank_line('Date and place')

newpage(); heading('APPROVAL SHEET', 1)
para('The project report “Credit Card Fraud Detection Using Machine Learning” is submitted for examination and approval by the following students:')
table(['No.', 'Student name', 'Roll / enrollment number'], [[i, '____________________________', '________________________'] for i in range(1, 5)])
blank_line('Examiner 1: name, signature, date')
blank_line('Examiner 2: name, signature, date')
blank_line('Project guide: name, signature, date')
blank_line('Place')

newpage(); heading('DECLARATION', 1)
para('We declare that this report describes the credit-card fraud detection project submitted with it. We have acknowledged the public dataset, external papers, and software documentation used in preparing the work. The experimental results reported here were taken from the project’s saved outputs, and the methods and limitations have been stated as implemented. The undersigned students accept responsibility for checking the final document and adding institutional details before submission.')
para('Student signatures:')
table(['No.', 'Student name', 'Signature', 'Date'], [[i, '___________________', '___________________', '____________'] for i in range(1, 5)])
blank_line('Place')

newpage(); heading('ACKNOWLEDGEMENT', 1)
para('We thank our project guide for reviewing the problem statement, experimental design, and presentation of the results. We also thank the department and institution for supporting this work. The public ULB / Worldline dataset, made accessible through Kaggle and the TensorFlow tutorial mirror, enabled the reproducible experiments in this report. The open-source Python ecosystem provided the tools used for analysis, model development, evaluation, and the demonstration interface.')
para('Guide, department, institution, and student names may be added after the report is reviewed.')
blank_line('Project guide')
blank_line('Department / institution')

newpage(); heading('ABSTRACT', 1)
para('Fraud detection is a difficult classification problem because fraudulent transactions are rare. This project develops and evaluates a machine learning workflow on the public ULB / Worldline credit-card dataset. The original 284,807 records contain 492 fraud cases. After removing 1,081 exact duplicates, the study uses 283,726 distinct records and 473 fraud cases. The data is divided into stratified 60% training, 20% validation, and 20% test partitions. Preprocessing, feature selection, and optional SMOTE resampling are fitted inside training pipelines to prevent evaluation leakage.')
para('Logistic Regression, Decision Tree, and Random Forest are each compared with and without SMOTE, both before and after bounded hyperparameter tuning. Validation average precision chooses the deployment configuration; validation F1 chooses its operating threshold. The selected Random Forest baseline does not use SMOTE. With the frozen threshold of 0.259462, it achieves 0.8795 precision, 0.7684 recall, 0.8202 F1, 0.9538 ROC-AUC, and 0.7792 average precision on 56,746 test transactions. It detects 73 of 95 frauds, misses 22, and creates 10 false alerts. A Streamlit application performs single and CSV batch inference using the saved pipeline.')
para('These measurements describe an educational benchmark. The PCA variables are anonymous, the dataset spans only two days in 2013, and the output scores have not been independently calibrated as real-world fraud probabilities. The project demonstrates careful handling of class imbalance and an end-to-end path from dataset audit to usable inference.')
para('Keywords: credit-card fraud, imbalanced classification, SMOTE, Random Forest, average precision, Streamlit.', 'Small')

newpage(); heading('CONTENTS', 1)
para('Update this table in Microsoft Word after editing the report: right-click it and select “Update Field”, then “Update entire table”.', 'Small')
field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
newpage(); heading('LIST OF FIGURES', 1)
para('Figure captions and their locations are included throughout the report. Update field references in Word if page numbers change.', 'Small')
figure_list_placeholder = doc.add_paragraph()

chapter(1, 'Introduction')
heading('1.1 Background', 2)
para('Electronic payments create large streams of transactions in which fraudulent activity forms a very small minority. A detector must identify suspicious records while avoiding excessive false alerts, which consume review resources and inconvenience legitimate customers. In the selected public benchmark, fraud appears in only about 0.173% of the raw data. A system that labels every transaction legitimate can therefore obtain about 99.83% accuracy while detecting no fraud. Accuracy alone is insufficient for this problem.')
para('The project addresses this imbalance by comparing three common supervised classifiers under a controlled workflow. The goal is a transparent educational comparison with a functional demonstration interface and recorded results. The model output is a screening signal for further review, not a confirmed finding of fraud.')
heading('1.2 Problem statement', 2)
para('Given a transaction represented by Time, Amount, and 28 anonymized PCA components, predict whether its label is legitimate (Class 0) or fraudulent (Class 1). The principal challenges are extreme class imbalance, preventing preprocessing and sampling leakage, choosing a useful operating threshold, and presenting model scores responsibly. The application must accept the same numeric feature schema that the trained model expects.')
heading('1.3 Objectives', 2)
for x in ['Audit and clean the original public dataset, preserving genuine minority examples and recording provenance.', 'Compare Logistic Regression, Decision Tree, and Random Forest with and without SMOTE.', 'Use validation data for model and threshold selection, then evaluate the frozen choice on untouched test data.', 'Report accuracy, precision, recall, F1, ROC-AUC, average precision, PR-AUC, and confusion counts.', 'Deliver an executable notebook, saved inference artifacts, and a Streamlit interface for single and batch scoring.']:
    bullet(x)
heading('1.4 Scope and expected use', 2)
para('The system is designed around the ULB / Worldline benchmark and a classroom demonstration. It scores numerical feature vectors supplied from the dataset or a compatible upstream transformation. It does not accept a card number, merchant name, or customer profile as a substitute for the anonymized PCA variables. The application displays an alert at a validation-selected threshold and allows a reviewer to download the results. Financial decisions, production monitoring, and direct card-system integration are outside the implemented scope.')
heading('1.5 Organization of the report', 2)
para('Chapter 2 reviews the data and methods relevant to imbalanced classification. Chapter 3 presents dataset handling, system design, and training methodology. Chapter 4 records implementation, results, interface behavior, and verification. Chapter 5 gives conclusions, limitations, and directions for further work. Appendices list feature selection, software and artifacts, and concise reproduction steps.')

chapter(2, 'Literature review and existing approaches')
heading('2.1 Dataset and benchmark context', 2)
para('The public ULB / Worldline dataset contains European cardholder transactions recorded over two days in September 2013. Its 28 variables V1–V28 are PCA-transformed and anonymized; Time is seconds since the first recorded transaction and Amount is transaction value. Class is the binary outcome. The Kaggle listing is the dataset source, and TensorFlow’s imbalanced-classification tutorial hosts the CSV mirror used by this project [1, 2]. The anonymization protects source information but prevents ordinary card fields from being transformed by this project.')
heading('2.2 Existing detection strategies', 2)
para('A fixed rule can flag a transaction above an amount threshold or outside a permitted pattern, but the exploratory results here show that fraud occurs at both low and high amounts. Statistical and machine learning models learn multivariable decision boundaries from labelled examples. Logistic Regression offers a regularized linear baseline; a Decision Tree makes nonlinear splits; and a Random Forest averages many randomized trees to reduce variance. None of these algorithms by itself solves the rarity of fraud or guarantees stable performance on future transactions.')
heading('2.3 Class imbalance and synthetic sampling', 2)
para('Chawla and colleagues introduced SMOTE to generate synthetic minority-class observations by interpolating among nearby minority samples [3]. It can improve detection sensitivity, but generated points can overlap legitimate patterns or be unrealistic. Imbalanced-learn explicitly warns that resampling before a train/test split leaks information and gives misleading evaluation [4]. Accordingly, this project performs any SMOTE operation only inside the training pipeline and leaves the validation and test partitions at their natural prevalence.')
heading('2.4 Evaluation measures', 2)
para('Precision asks how many alerts are actual fraud; recall asks how many actual frauds are caught. F1 combines them at a chosen threshold. The confusion matrix records true negatives, false positives, false negatives, and true positives. ROC-AUC summarizes ranking across true- and false-positive rates, while average precision summarizes the precision–recall curve using recall-weighted increments [5]. Trapezoidal PR-AUC is reported separately because its interpolation differs from average precision. Under extreme imbalance, PR measures and the operational alert counts convey information that high accuracy can obscure.')
heading('2.5 Gap addressed by this project', 2)
para('The project connects a reproducible benchmark comparison to a working inference path. It audits duplicates, fits all learned transformations on training data, compares six configurations before and after tuning, freezes model and threshold using validation data, and records held-out test metrics. It then saves a complete pipeline and uses it in a Streamlit interface, so the documented preprocessing and the application inference path agree.')

chapter(3, 'Proposed system and methodology')
heading('3.1 Data collection and provenance', 2)
para('The training script reads data/creditcard.csv after download from the public TensorFlow mirror. The downloader checks the expected raw record and fraud counts and records a SHA-256 digest. The specific CSV used for the reported run has SHA-256 76274b691b16a6c49d3f159c883398e03ccd6d1ee12d9d8ee38f4b4b98551a89. If the mirror is unavailable, the original Kaggle archive can be downloaded manually and creditcard.csv placed in the data directory. No synthetic replacement dataset is used for the reported model results.')
table(['Attribute', 'Meaning / treatment'], [
    ('Time', 'Seconds since first transaction; nonnegative numeric input.'),
    ('V1–V28', 'Anonymized PCA components; numeric values may be positive or negative.'),
    ('Amount', 'Recorded transaction amount; nonnegative numeric input.'),
    ('Class', 'Target only: 0 legitimate, 1 fraud; excluded during prediction.')])
heading('3.2 Audit, cleaning and split', 2)
para('The loader checks required columns, numeric types, target labels, finite inputs, and valid nonnegative Time and Amount. It found no missing values in the original CSV. The audit removes 1,081 exact full-row duplicates before partitioning, including 19 fraud duplicates. Removing exact copies before splitting prevents the same transaction from appearing on both sides of an evaluation boundary. Identical predictors with conflicting labels cause a validation error rather than being silently assigned to partitions. Outliers are retained because an extreme transaction can be a genuine fraud signal.')
table(['Dataset stage', 'All transactions', 'Fraud cases', 'Fraud share'], [
    ('Original CSV', '284,807', '492', '0.1727%'),
    ('After exact deduplication', '283,726', '473', '0.1667%'),
    ('Training', '170,235', '284', '0.1668%'),
    ('Validation', '56,745', '94', '0.1657%'),
    ('Test', '56,746', '95', '0.1674%')])
para('The cleaned records are stratified into 60% training, 20% validation, and 20% test with random seed 42. The training partition supports model fitting and design-oriented exploratory analysis. Validation guides model and threshold selection. Test labels are used only after those decisions are frozen. Original row indices are retained for partition audits.')
image(FIG / 'class_imbalance.png', 'Fraud is rare in the cleaned training partition.', 5.25)
heading('3.3 Exploratory data analysis', 2)
para('Training-only plots examine class frequencies, Time and Amount distributions, correlations, PCA component distributions, and Time versus Amount. Legitimate transactions have median Amount 22.05, while fraudulent transactions have median Amount 11.86 in the training data. This difference does not imply a simple low-amount fraud rule: the classes overlap. For visualization, a scatter plot samples at most 3,000 legitimate records together with all training frauds; that plotted subset does not represent natural prevalence.')
para('The strongest absolute linear correlations with the target among training features are V17 (0.320), V14 (0.299), V12 (0.251), and V10 (0.208). Correlation is descriptive and cannot establish a cause, especially with anonymized PCA components. Outliers are shown or visually compressed in plots but remain in the model input.')
image(FIG / 'amount_time_distributions.png', 'Training Time and Amount distributions by class.', 5.85)
image(FIG / 'correlation_heatmap.png', 'Training-only feature and target correlations.', 5.15)
heading('3.4 Preprocessing and feature selection', 2)
para('The preprocessing pipeline first uses median imputation, although the audited original data has no missing values. It then applies RobustScaler to Time and Amount, which scales with the median and interquartile range, and StandardScaler to V1–V28. Scaling is needed for the linear model and for the neighbor distances used by SMOTE. These operations are fitted only on the training partition or, in cross-validation, on each training fold. Categorical encoding is unnecessary because all 30 predictors are numeric.')
para('A custom selector estimates continuous mutual information between each scaled predictor and Class on up to 60,000 stratified training rows. The predetermined top 20 of 30 predictors are retained. In the saved full-training fit, this sample contains 100 frauds and selects Time, V1–V12, V14, V16–V19, V21, and V27. The precise columns can differ across cross-validation folds because the selector is fitted anew within each fold. A separate training-only Random Forest importance chart is descriptive and does not alter feature selection.')
image(FIG / 'feature_selection.png', 'Mutual-information ranking from training data; the top 20 are selected.', 5.45)
image(FIG / 'feature_importance.png', 'Training-only Random Forest impurity importance for selected features.', 5.45)
heading('3.5 Classification and imbalance handling', 2)
para('The baseline classifiers are Logistic Regression (C=1.0, lbfgs, maximum 1,000 iterations), Decision Tree (maximum depth 10, minimum leaf size 2), and Random Forest (80 trees, maximum depth 14, minimum leaf size 2, 70% sample fraction). The same feature pipeline is used with each algorithm. Each baseline is fitted once without resampling and once with SMOTE, producing six baseline experiments. SMOTE uses five neighbors and targets a minority-to-majority ratio of 0.10. The sampling step is after scaling and selection in an imbalanced-learn pipeline; it is active during fitting and skipped during prediction.')
table(['Classifier', 'Strength in this experiment', 'Main limitation'], [
    ('Logistic Regression', 'Efficient regularized baseline with score output', 'Linear decision boundary can miss interactions'),
    ('Decision Tree', 'Nonlinear, simple split rules', 'High variance with rare positive cases'),
    ('Random Forest', 'Averages trees and captures nonlinear patterns', 'More computational cost; indirect interpretation')])
heading('3.6 Hyperparameter search and selection rule', 2)
para('All six configurations undergo a bounded RandomizedSearchCV: three sampled parameter combinations and three stratified folds, scored by average precision. Search runs on a natural-prevalence sample of at most 60,000 training records containing 100 frauds. Each best parameter set is refitted on all 170,235 training records. Fold-local preprocessing, feature selection, and any SMOTE operation are recomputed inside the pipeline. Each search fold has about 33–34 validation frauds, so small differences in cross-validation scores are uncertain.')
para('The deployment configuration is the baseline or tuned experiment with the highest validation average precision, breaking ties by validation F1 and then experiment name. For that one configuration, the operating threshold is chosen to maximize validation F1; exact ties choose the highest threshold. The selected model and threshold are written to a frozen-selection file before any test evaluation. The 0.5 threshold remains useful for fair comparisons, while the frozen threshold defines the final operating result.')
table(['Stage', 'Partition / rule', 'Leakage control'], [
    ('Fit and tune', 'Training; 3-fold stratified search', 'All learned steps stay within each training fold'),
    ('Model selection', 'Independent validation; maximum AP', 'No test labels consulted'),
    ('Threshold choice', 'Validation; maximum F1', 'One threshold frozen before test'),
    ('Final evaluation', 'Untouched test partition', 'No retuning from test outcomes')])
heading('3.7 Proposed architecture and modules', 2)
para('The workflow is: validated CSV → deduplication → stratified partitions → imputation and scaling → mutual-information selection → optional SMOTE during fitting → classifier → saved pipeline and metadata → Streamlit single or batch prediction. At inference, the application validates 30 inputs and feeds the raw compatible values through the complete saved pipeline exactly once. The separate saved scaler is used for consistency checks, not for an extra scaling pass.')
table(['Module / file', 'Responsibility'], [
    ('src/data.py', 'Load, validate, audit, deduplicate and stratify the CSV.'),
    ('src/preprocessing.py', 'Median imputation, scaling and mutual-information selector.'),
    ('src/training.py', 'Six baselines, tuning, selection, test evaluation and saved artifacts.'),
    ('src/evaluation.py', 'Metrics, threshold selection, ROC/PR and confusion plots.'),
    ('src/predict.py', 'Artifact checks, transaction validation and inference.'),
    ('src/reporting.py', 'Data analysis figures and generated findings.'),
    ('app.py', 'Streamlit single and batch demonstration interface.'),
    ('notebooks/', 'Eight ordered project components and executed explanations.')])
heading('3.8 Software and hardware requirements', 2)
para('The project is specified for Python 3.12 and pinned packages in requirements.txt. Core data and model libraries are NumPy 2.2.6, pandas 2.3.3, SciPy 1.16.2, scikit-learn 1.7.2, imbalanced-learn 0.14.0, and joblib 1.5.2. Matplotlib 3.10.7 and seaborn 0.13.2 generate figures; Streamlit 1.50.0 serves the interface; Jupyter packages support the notebook; pytest 8.4.2 runs checks. No dedicated GPU is required by these algorithms. Sufficient disk space is needed for the approximately 151 MB source CSV, artifacts, notebook, and figures. Training time depends on CPU and selected search budget.')

chapter(4, 'Implementation, results and discussion')
heading('4.1 Implementation sequence', 2)
para('The reproducible workflow begins by installing pinned dependencies, downloading and auditing the real CSV, and preparing training-only exploratory summaries. The training script fits six baseline pipelines, searches a limited number of parameter candidates for all six configurations, refits them, then records validation and frozen test results. The notebook reads the saved figures and tables and can rerun training when requested. The application reads only saved deployment artifacts and does not need a training operation at startup.')
table(['Milestone', 'Completed evidence'], [
    ('Data and setup', '284,807 original rows; source digest recorded.'),
    ('Analysis and preprocessing', '1,081 duplicates removed; EDA and feature rankings saved.'),
    ('Baselines', 'Six fitted configurations and validation metrics.'),
    ('Tuning and selection', 'Six searches, each with 3 candidates × 3 folds; model and threshold frozen.'),
    ('Notebook and deployment', '28 cells execute without errors; joblib pipeline and Streamlit app saved.'),
    ('Verification', '23 pytest checks pass; browser interactions and local HTTP startup verified.')])
heading('4.2 Evaluation definitions', 2)
table(['Measure', 'Meaning for this project'], [
    ('Accuracy', 'Correct predictions divided by all transactions; dominated by Class 0.'),
    ('Precision', 'True fraud alerts divided by all fraud alerts; relates to alert workload.'),
    ('Recall', 'Detected frauds divided by all actual frauds; relates to missed fraud.'),
    ('F1', 'Harmonic mean of precision and recall at a chosen threshold.'),
    ('ROC-AUC', 'Ranking performance over true-positive and false-positive rates.'),
    ('Average precision', 'Recall-weighted precision; primary tuning and selection measure.'),
    ('PR-AUC', 'Trapezoidal precision–recall area; reported separately from AP.'),
    ('TN / FP / FN / TP', 'Correct legitimate / false alert / missed fraud / detected fraud.')])
para('A constant always-legitimate classifier illustrates why accuracy is misleading. On the deduplicated test set it labels all 56,746 records legitimate, attaining roughly 99.83% accuracy but zero fraud recall. For ranking, its AP equals test prevalence, approximately 0.00167. The main tables therefore emphasize AP, precision, recall, and confusion counts.')
heading('4.3 Validation comparison', 2)
para('Table 4.3 uses the saved validation results at the common threshold of 0.5. “None” indicates no SMOTE. The baseline Random Forest without SMOTE has the highest validation AP, 0.8659, although it was not the best in every other measure. The tuned Random Forest without SMOTE has AP 0.8579. The validation selection rule therefore retains the untuned baseline for deployment.')
rows=[]
for r in comparison:
    if r['partition']=='validation' and r['stage'] in ('baseline','tuned'):
        rows.append((r['algorithm'].replace('Logistic Regression','Logistic Reg.').replace('Decision Tree','Decision Tree').replace('Random Forest','Random Forest'), r['imbalance'],r['stage'],metric(r['precision'],3),metric(r['recall'],3),metric(r['f1'],3),metric(r['average_precision'],3)))
table(['Algorithm','SMOTE','Stage','Precision','Recall','F1','AP'], rows)
para('At threshold 0.5, SMOTE raises Logistic Regression recall from 0.660 to 0.894 in the baseline comparison, but its precision falls from 0.899 to 0.418 and false alerts rise from 7 to 117. For the baseline Random Forest, SMOTE raises recall from 0.830 to 0.851 but lowers AP from 0.8659 to 0.8326. These observed tradeoffs explain why sampling should be evaluated, rather than assumed to improve the final choice.')
image(FIG / 'validation_pr_curves.png', 'Precision–recall curves for all validation configurations.', 5.75)
heading('4.4 Tuning findings', 2)
para('The bounded search selected candidate parameters by cross-validation AP on 60,000 training rows. A better cross-validation score does not guarantee a better score on the separate validation partition. The untuned Random Forest without SMOTE remained the selected configuration. Decision Tree without SMOTE improved validation AP from 0.7185 to 0.8342 after tuning, while Decision Tree with SMOTE fell from 0.5992 to 0.5509. These changes show that synthetic sampling and parameter choices interact with the classifier.')
table(['Algorithm','SMOTE','CV AP','CV SD','Validation AP'], [
    (r['algorithm'].replace('Logistic Regression','Logistic Reg.'),r['imbalance'],metric(r['cv_average_precision'],3),metric(r['cv_ap_std'],3),metric(next(c['average_precision'] for c in comparison if c['experiment']==r['experiment'] and c['partition']=='validation'),3)) for r in tuning])
para('With only about 33–34 frauds in each search validation fold, the CV standard deviations are material. The 3 × 3 search was selected to bound computation for a mini-project; it cannot establish global optimality. Each tuned model was nonetheless refitted on the full training partition before validation comparison.')
heading('4.5 Frozen model and operating threshold', 2)
para('The selected configuration is Random Forest | None | baseline, chosen using validation AP 0.8658818. The saved model is the original training-only fit, without refitting on validation or test data. The threshold 0.2594623 maximized validation F1 and was frozen before the test set was scored. A lower threshold can improve fraud coverage at the cost of additional review alerts. A production threshold would require known investigation capacity, missed-fraud losses, and customer-impact costs, which are not available in this benchmark.')
table(['Selection item','Recorded value'], [
    ('Model','Random Forest, baseline parameters, no SMOTE'),
    ('Validation AP','0.8658818'),
    ('Threshold selection','Maximum validation F1'),
    ('Saved operating threshold','0.2594623'),
    ('Test used for model or threshold choice','No')])
image(FIG / 'final_threshold_and_confusion.png', 'Frozen threshold analysis and final test confusion matrix.', 6.0)
heading('4.6 Held-out test performance', 2)
para('The frozen model and threshold were evaluated on 56,746 previously untouched records: 56,651 legitimate and 95 fraudulent transactions. It correctly classifies 56,641 legitimate records and 73 fraud cases, while producing 10 false alerts and missing 22 frauds. The final measurements are summarized below. These figures describe one particular public-data split, not an estimate for all payment systems.')
final = meta['final_test_metrics']
table(['Metric','Value'],[(k,v) for k,v in [
    ('Accuracy',metric(final['accuracy'],6)),('Precision',metric(final['precision'],4)),('Recall',metric(final['recall'],4)),
    ('F1',metric(final['f1'],4)),('ROC-AUC',metric(final['roc_auc'],4)),('Average precision',metric(final['average_precision'],4)),
    ('Trapezoidal PR-AUC',metric(final['pr_auc'],4)),('TN / FP / FN / TP','56,641 / 10 / 22 / 73')]])
para('For the same model at the default 0.5 threshold, the test set has 3 false alerts and 26 missed frauds. The frozen 0.259462 threshold detects four more frauds and adds seven false alerts. Test F1 decreases slightly from 0.8263 to 0.8202; that outcome was not used to revisit the validation-selected threshold. ROC-AUC and AP remain unchanged when only the final decision threshold changes because the underlying scores and their ordering are identical.')
image(FIG / 'test_pr_curves.png', 'Test precision–recall curves for the frozen set of candidate models.', 5.75)
image(FIG / 'test_roc_curves.png', 'Test ROC curves for the frozen candidate models.', 5.75)
heading('4.7 Error analysis and interpretation', 2)
para('False negatives are the 22 frauds the selected operating rule misses. In a real setting, they could lead to monetary loss and delayed response. False positives are 10 legitimate transactions that would enter manual review and could inconvenience customers. These counts are more actionable than overall accuracy. The baseline Decision Tree shows a sizeable training-to-validation AP gap (0.9068 versus 0.7185), consistent with overfitting risk; the gap alone is not proof because rare fraud samples vary between partitions. The forest reduces variance by averaging trees but still has a training-to-validation gap and an uncertain test estimate.')
para('Mutual-information ranking identifies statistical association, not a causal fraud mechanism. The top saved forest impurity importances include V17, V12, V14, V11, and V16. The original meanings of those PCA dimensions are unavailable, so the charts cannot be translated into explanations such as merchant type or cardholder behavior. The displayed model score comes from predict_proba and is not independently calibrated as a real fraud probability.')
heading('4.8 Streamlit application', 2)
para('The application, titled “Fraud Signal Studio”, offers two tabs. The single-transaction tab accepts Time, Amount, and all 28 PCA components, and can preload legitimate or fraudulent examples taken from held-out records. After submission it shows the model score, the saved decision threshold, an alert label, an input-profile chart, and global forest importance. The batch tab accepts a CSV with all 30 numeric inputs, scores every row, summarizes alert counts and score distribution, and offers a scored CSV download. A blank schema template and an example batch are available. If Class is present in an uploaded example, it is ignored for prediction.')
para('Input checks reject missing or duplicate required columns, nonnumeric or nonfinite values, negative Time or Amount, empty batches, and invalid thresholds. The app checks that the saved model, feature list, scaler, metadata, and scikit-learn version agree. These checks prevent silently scoring a row with the wrong schema or applying a separately saved scaler twice. The four screenshots below were captured from the running Streamlit application using real saved artifacts; they are not mockups.')
image(ASSETS / 'app_single_input.png', 'Running application: single-transaction input form and model notes.', 5.5)
image(ASSETS / 'app_fraud_result.png', 'Running application: held-out fraud example flagged for review.', 5.5)
image(ASSETS / 'app_batch_results.png', 'Running application: scored CSV batch and summary.', 5.5)
image(ASSETS / 'app_invalid_csv.png', 'Running application: readable error for an incomplete CSV schema.', 5.5)
heading('4.9 Testing and verification', 2)
para('The test suite covers schema validation, data cleaning and splitting, fold-local preprocessing, metric calculations, artifact consistency, inference, notebook structure, and Streamlit interactions. On 8 October 2026, a fresh run of python -m pytest -q finished with 23 passed checks. The project’s executed notebook record reports all 28 cells completed with zero errors. An HTTP startup check previously recorded status 200 for the health and page endpoints. Browser verification for this report loaded a real held-out fraud example, scored an uploaded example CSV, and checked the invalid-CSV message. The screenshots serve as visual evidence of the running interface.')
table(['Verification','Outcome / evidence'], [
    ('Unit and integration suite','23 passed; .cache/report-verification.xml'),
    ('Notebook execution','28 cells, zero errors; results/notebook_execution.json'),
    ('HTTP startup','Health 200, page 200; results/server_verification.json'),
    ('Browser demonstration','Fraud example, CSV upload and invalid CSV checked; reports/assets/')])
heading('4.10 Reproducibility and artifacts', 2)
para('The project records a fixed random seed of 42, pinned package versions, the original CSV SHA-256 digest, split counts, parameter-search output, selected model, threshold, and results. The serialized best_model.joblib is a complete inference pipeline; scaler.joblib, preprocessor.joblib, selected_features.joblib, and metadata.json support inspection and consistency checks. The full data and code are local to the project directory. Training, notebook execution, and UI startup can be reproduced with the commands in Appendix C, assuming the public dataset is available and the pinned environment installs successfully.')

chapter(5, 'Conclusion and future scope')
heading('5.1 Conclusion', 2)
para('The project demonstrates a complete machine learning workflow for highly imbalanced credit-card transaction classification. It begins with the real public ULB / Worldline CSV, removes exact duplicates, uses disjoint stratified partitions, fits preprocessing and optional SMOTE inside training pipelines, compares three algorithms under six sampling configurations, and evaluates the final frozen choice on held-out data. The Streamlit application successfully uses the saved model for compatible single and batch inputs.')
para('The Random Forest baseline without SMOTE had the highest validation AP and was selected with a validation-chosen operating threshold. On the test split, it found 73 of 95 frauds with 10 false alerts. Its precision of 0.8795 and recall of 0.7684 summarize the tradeoff at that threshold. The result supports the value of comparing measured precision and recall instead of relying on accuracy or assuming that oversampling must help.')
heading('5.2 Limitations', 2)
for x in [
    'The dataset covers only two days in September 2013; present-day transaction behavior and fraud tactics can differ.',
    'PCA variables are anonymized and the original transformation is unavailable, so ordinary card details cannot be directly scored.',
    'The random stratified split does not test performance on future time periods or across cardholders; card and customer identifiers are unavailable for grouping.',
    'The test set contains only 95 frauds. Estimates can vary appreciably with a different split or operating environment.',
    'Exact deduplication changes the record counts relative to published raw-data benchmarks.',
    'Mutual information ranks features separately and may miss useful interactions; the top-20 count was fixed for this project, not proven optimal.',
    'The search evaluates only three candidates per configuration, and SMOTE may generate unrealistic points in an anonymized feature space.',
    'Reported scores are uncalibrated, and the validation-F1 threshold is not a substitute for bank-specific review costs and capacity.'
]: bullet(x)
heading('5.3 Future work', 2)
for x in [
    'Evaluate a future-time holdout and, if identifiers become available, split by cardholder to measure generalization more realistically.',
    'Compare class weighting, balanced forests, and calibrated probabilities with the present SMOTE approach.',
    'Run larger or nested cross-validation and feature-count ablation to estimate uncertainty and model-selection sensitivity.',
    'Choose thresholds from actual investigation capacity, financial losses, and customer-impact costs.',
    'Monitor data and score drift, retrain on more recent data, and validate alerts with human analysts.',
    'Integrate a compatible upstream PCA transformation before considering raw transaction entry in a real payment workflow.'
]: bullet(x)

newpage(); heading('REFERENCES', 1)
references = [
    '[1] Machine Learning Group, Université Libre de Bruxelles. “Credit Card Fraud Detection” (ULB / Worldline dataset), Kaggle. https://www.kaggle.com/datasets/mlg-ulb/creditcardfraud',
    '[2] TensorFlow. “Classification on imbalanced data.” https://www.tensorflow.org/tutorials/structured_data/imbalanced_data',
    '[3] N. V. Chawla, K. W. Bowyer, L. O. Hall and W. P. Kegelmeyer. “SMOTE: Synthetic Minority Over-sampling Technique.” Journal of Artificial Intelligence Research, vol. 16, pp. 321–357, 2002. https://doi.org/10.1613/jair.953',
    '[4] imbalanced-learn documentation. “Common pitfalls and recommended practices.” https://imbalanced-learn.org/stable/common_pitfalls.html',
    '[5] scikit-learn 1.7 documentation. “average_precision_score.” https://scikit-learn.org/1.7/modules/generated/sklearn.metrics.average_precision_score.html',
    '[6] scikit-learn documentation. “mutual_info_classif.” https://scikit-learn.org/stable/modules/generated/sklearn.feature_selection.mutual_info_classif.html',
    '[7] Streamlit documentation. “Get started.” https://docs.streamlit.io/get-started',
    '[8] Project repository artifacts: README.md; src/; notebooks/credit_card_fraud_detection.ipynb; results/; models/ (local submitted project).'
]
for ref in references: para(ref)

newpage(); heading('APPENDIX A  |  FEATURE SELECTION', 1)
para('The following values are the saved full-training mutual-information ranking. The selector used a stratified 60,000-row sample with 100 frauds. The final 20-feature set is recorded in models/metadata.json. Fold-level sets may differ.')
table(['Rank','Feature','Mutual information','Selected'], [(i,r['feature'],f"{float(r['mutual_information']):.6f}",'Yes' if r['selected']=='True' else 'No') for i,r in enumerate(feature,1)])
para('Selected raw inputs for the saved model: ' + ', '.join(meta['selected_features']) + '.')

newpage(); heading('APPENDIX B  |  PROJECT FILES AND CONFIGURATION', 1)
table(['Path','Purpose'], [
    ('data/creditcard.csv','Original public CSV used for training.'),('data/creditcard.source.json','Dataset source, counts and SHA-256 digest.'),
    ('notebooks/credit_card_fraud_detection.ipynb','Eight ordered ML workflow sections.'),('src/','Reusable data, preprocessing, training, evaluation and inference modules.'),
    ('scripts/','Download, preparation, training, notebook and server commands.'),('figures/','Saved training, validation and test plots.'),
    ('results/','Audit, model comparison, tuning, selection, checks and examples.'),('models/','Serialized pipeline and supporting deployment metadata.'),
    ('app.py','Single and batch Streamlit interface.'),('tests/','Automated test suite.'),('requirements.txt','Pinned Python dependencies.')])
para('Reported training configuration: seed 42; 60/20/20 stratified split; top 20 of 30 features; at most 60,000 rows for selector ranking; SMOTE target ratio 0.10; three stratified CV folds; three tuning candidates per configuration on at most 60,000 rows; two jobs. More exhaustive training is possible through scripts/train.py command-line options, but this report records the saved default run.')

newpage(); heading('APPENDIX C  |  REPRODUCTION AND DEMONSTRATION', 1)
para('From the project root on Python 3.12, run the following commands. The local virtual environment may be used in place of the python command:')
for command in ['python -m venv .venv', '.\\.venv\\Scripts\\Activate.ps1', 'python -m pip install -r requirements.txt', 'python scripts/download_data.py', 'python scripts/train.py', 'python -m pytest -q', 'python scripts/execute_notebook.py', 'python -m streamlit run app.py']:
    para(command, 'Small')
para('The app opens at http://localhost:8501 by default. For a single demo, choose “Fraud example” in the Single transaction tab, then select “Predict transaction”. For batch scoring, use “Score the example batch” or upload a CSV with Time, V1–V28 and Amount. Class is optional and ignored for inference. The score is uncalibrated; it is intended to help prioritize review.')

newpage(); heading('APPENDIX D  |  SHORT VIVA NOTES', 1)
table(['Question','Answer'], [
    ('Why is accuracy insufficient?','Always legitimate achieves about 99.83% accuracy but finds no fraud.'),
    ('What is SMOTE?','Synthetic interpolation between minority training points.'),
    ('Where is SMOTE applied?','Only inside training folds after scaling and selection.'),
    ('Why use AP?','It summarizes precision against recall under severe imbalance.'),
    ('How was the final model chosen?','Maximum validation AP among baseline and tuned models.'),
    ('How was threshold chosen?','Maximum validation F1, then frozen before test.'),
    ('What is a false negative?','A fraud classified as legitimate; 22 occurred on the test set.'),
    ('Can normal card details be entered?','No; the system needs compatible Time, V1–V28 and Amount values.'),
    ('Are scores real fraud probabilities?','No independent calibration was performed.'),
    ('What is the main practical limitation?','Old, short, anonymized data without temporal or cardholder validation.')])

figure_list_placeholder.text = '\n'.join(f'Figure {i}. {title}' for i,title in enumerate(figure_index,1))
doc.core_properties.title = 'Credit Card Fraud Detection Using Machine Learning'
doc.core_properties.subject = 'Project report based on saved ULB credit-card fraud detection experiments'
doc.core_properties.keywords = 'credit card fraud; machine learning; SMOTE; Random Forest; Streamlit'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
dest = OUT / 'Credit_Card_Fraud_Detection_Project_Report.docx'
doc.save(dest)
print(f'Saved {dest} ({dest.stat().st_size:,} bytes), {len(figure_index)} figures')
